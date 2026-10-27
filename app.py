
import io
import os
import re
import json
import random
import sqlite3
import tempfile
from pathlib import Path
from datetime import datetime

import streamlit as st

# ============================================================
# EDUVERSE AI LAB
# Single-file study platform: notes, audio, speaking, quizzes,
# flashcards, documents, analytics, and resource library.
# ============================================================

APP_DIR = Path(__file__).resolve().parent
DB_PATH = APP_DIR / "eduverse.db"

st.set_page_config(
    page_title="EduVerse AI Lab",
    page_icon="🎓",
    layout="wide",
    initial_sidebar_state="expanded",
)

# --------------------------- DATABASE -------------------------


def get_db():
    conn = sqlite3.connect(DB_PATH, check_same_thread=False)

    # Expected schemas for the current version
    expected_tables = {
        "activity": {
            "id", "kind", "title", "score", "created"
        },
        "resources": {
            "id", "title", "kind", "body", "created"
        },
    }

    for table_name, expected_columns in expected_tables.items():
        existing = conn.execute(
            f"PRAGMA table_info({table_name})"
        ).fetchall()

        if existing:
            actual_columns = {row[1] for row in existing}

            # Back up an incompatible old table instead of deleting it.
            if not expected_columns.issubset(actual_columns):
                backup_name = (
                    f"{table_name}_legacy_"
                    + datetime.now().strftime("%Y%m%d%H%M%S")
                )

                conn.execute(
                    f"ALTER TABLE {table_name} RENAME TO {backup_name}"
                )

    conn.execute("""
        CREATE TABLE IF NOT EXISTS activity (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            kind TEXT NOT NULL,
            title TEXT NOT NULL,
            score REAL,
            created TEXT NOT NULL
        )
    """)

    conn.execute("""
        CREATE TABLE IF NOT EXISTS resources (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            title TEXT NOT NULL,
            kind TEXT NOT NULL,
            body TEXT NOT NULL,
            created TEXT NOT NULL
        )
    """)

    conn.commit()
    return conn



db = get_db()


def log_activity(kind, title, score=None):
    db.execute(
        "INSERT INTO activity(kind,title,score,created) VALUES(?,?,?,?)",
        (kind, title[:200], score, datetime.now().isoformat(timespec="seconds")),
    )
    db.commit()


def save_resource(title, kind, body):
    db.execute(
        "INSERT INTO resources(title,kind,body,created) VALUES(?,?,?,?)",
        (
            title[:200],
            kind[:100],
            body,
            datetime.now().isoformat(timespec="seconds"),
        ),
    )
    db.commit()


# ---------------------------- AI SETUP -----------------------

def get_api_key():
    key = os.getenv("OPENAI_API_KEY", "").strip()
    if key:
        return key

    try:
        return str(st.secrets.get("OPENAI_API_KEY", "")).strip()
    except Exception:
        return ""


def ai_request(prompt, system="You are a careful study assistant.", temperature=0.2):
    """Return (answer, error). No fake success when AI is unavailable."""
    key = get_api_key()

    if not key:
        return None, "No API key configured. The offline method will be used where available."

    try:
        from openai import OpenAI

        client = OpenAI(api_key=key)
        response = client.chat.completions.create(
            model=os.getenv("OPENAI_MODEL", "gpt-4o-mini"),
            messages=[
                {"role": "system", "content": system},
                {"role": "user", "content": prompt},
            ],
            temperature=temperature,
        )

        answer = response.choices[0].message.content
        if not answer or not answer.strip():
            return None, "The AI service returned an empty answer."

        return answer.strip(), None

    except Exception as exc:
        return None, f"AI request failed: {exc}"


# ------------------------- TEXT UTILITIES ---------------------

def normalize(text):
    return re.sub(r"\s+", " ", text or "").strip()


def split_sentences(text):
    parts = re.split(r"(?<=[.!?])\s+|\n+", (text or "").strip())
    output = []

    for part in parts:
        part = part.strip(" \t\r\n-•*")
        if len(part) >= 12:
            output.append(part)

    return output


def parse_json(text):
    if not text:
        return None

    try:
        return json.loads(text)
    except Exception:
        pass

    match = re.search(r"\[[\s\S]*\]", text)
    if match:
        try:
            return json.loads(match.group(0))
        except Exception:
            return None

    return None


def extract_terms(text, limit=20):
    words = re.findall(r"\b[A-Za-z][A-Za-z0-9_-]{3,}\b", text or "")
    stop = {
        "this", "that", "these", "those", "there", "their", "which",
        "where", "when", "what", "with", "from", "into", "have",
        "been", "being", "were", "will", "would", "could", "should",
        "about", "because", "while", "also", "than", "then", "them",
        "they", "your", "more", "some", "such", "only", "very",
        "using", "used", "each", "other", "both", "many", "most",
    }

    counts = {}
    for word in words:
        low = word.lower()
        if low not in stop:
            counts[low] = counts.get(low, 0) + 1

    return [
        word for word, _ in sorted(
            counts.items(), key=lambda item: (-item[1], item[0])
        )[:limit]
    ]


# --------------------- SOURCE-BASED NOTES ---------------------

def offline_notes(source, title):
    """
    Extracts actual source sentences. It does not invent missing
    definitions, examples, or explanations.
    """
    sentences = split_sentences(source)

    if not sentences:
        return (
            f"# {title}\n\n"
            "The supplied text did not contain enough readable sentences "
            "to create notes. Please review the transcript or document."
        )

    terms = extract_terms(source, 15)

    lines = [
        f"# {title}",
        "",
        "## Source-based overview",
        "These notes are extracted from the text you supplied. "
        "The offline method does not add facts that are absent from the source.",
        "",
        "## Important points from the source",
    ]

    for sentence in sentences[:25]:
        lines.append(f"- {sentence}")

    lines += ["", "## Terms appearing in the source"]

    if terms:
        lines.append(", ".join(terms))
    else:
        lines.append("No reliable terms could be extracted automatically.")

    lines += [
        "",
        "## Revision reminder",
        "Review the original transcript alongside these points. "
        "The offline method extracts content but does not infer missing explanations.",
    ]

    return "\n".join(lines)


def generate_notes(source, title):
    source = source.strip()

    if len(source) < 40:
        return None, "There is not enough transcript text to create useful notes."

    prompt = f"""
Create detailed study notes using ONLY the source text below.

Rules:
- Do not introduce a different chapter or unrelated topic.
- Do not invent facts, definitions, examples, formulas, or conclusions.
- Keep the original meaning and subject.
- If the source is unclear or incomplete, state that clearly.
- Do not claim that information is present when it is absent.

Use these sections where supported by the source:
1. Chapter/topic overview
2. Detailed notes organized by topics actually discussed
3. Important points
4. Definitions and examples explicitly present in the source
5. Formulas or processes mentioned in the source
6. Revision questions answerable from the source

Title: {title}

SOURCE TEXT:
--- START ---
{source[:30000]}
--- END ---
"""

    answer, error = ai_request(
        prompt,
        "You prepare faithful study notes. Source accuracy is more important than completeness.",
    )

    if answer:
        return answer, "AI-generated from the reviewed source text."

    return offline_notes(source, title), (
        "Offline source extraction used. "
        "For deeper explanations, configure an AI API key."
    )


# -------------------------- AUDIO --------------------------------

def transcribe_audio(upload):
    if upload is None:
        return None, "Record or upload an audio file first."

    audio_bytes = upload.getvalue()
    filename = getattr(upload, "name", "recording.wav")
    suffix = Path(filename).suffix.lower() or ".wav"

    # First attempt: local Whisper. It requires the package and FFmpeg.
    try:
        import whisper

        with tempfile.NamedTemporaryFile(delete=False, suffix=suffix) as temp:
            temp.write(audio_bytes)
            temp_path = temp.name

        try:
            model_name = os.getenv("WHISPER_MODEL", "base")
            model = whisper.load_model(model_name)
            result = model.transcribe(temp_path)
            transcript = (result.get("text") or "").strip()

            if transcript:
                return transcript, None

            return None, "No speech was detected. Try a clearer recording."

        finally:
            try:
                os.remove(temp_path)
            except OSError:
                pass

    except ImportError:
        pass

    except Exception as exc:
        whisper_error = str(exc)
    else:
        whisper_error = ""

    # Fallback: SpeechRecognition supports a limited set of formats.
    if suffix not in {".wav", ".flac", ".aiff", ".aif"}:
        message = (
            f"Could not transcribe {suffix} with the available fallback. "
            "Install Whisper and FFmpeg, or convert the audio to WAV."
        )
        if whisper_error:
            message += f" Whisper error: {whisper_error}"
        return None, message

    try:
        import speech_recognition as sr

        with tempfile.NamedTemporaryFile(delete=False, suffix=suffix) as temp:
            temp.write(audio_bytes)
            temp_path = temp.name

        try:
            recognizer = sr.Recognizer()

            with sr.AudioFile(temp_path) as source:
                audio = recognizer.record(source)

            try:
                transcript = recognizer.recognize_google(audio)
                if transcript.strip():
                    return transcript.strip(), None
                return None, "No speech was detected."

            except sr.UnknownValueError:
                return None, (
                    "The speech was unclear. Try recording in a quiet room "
                    "and speaking closer to the microphone."
                )

            except sr.RequestError as exc:
                return None, (
                    "The speech-recognition service is unavailable. "
                    f"Check your internet connection. Details: {exc}"
                )

        finally:
            try:
                os.remove(temp_path)
            except OSError:
                pass

    except ImportError:
        return None, (
            "No transcription engine is installed. Install "
            "SpeechRecognition or openai-whisper."
        )

    except Exception as exc:
        return None, f"Audio transcription failed: {exc}"


# ----------------------- DOCUMENT UTILITIES ---------------------

def extract_document(upload):
    filename = upload.name.lower()
    data = upload.getvalue()

    if filename.endswith((".txt", ".md", ".csv")):
        return data.decode("utf-8", errors="replace").strip()

    if filename.endswith(".pdf"):
        from pypdf import PdfReader

        reader = PdfReader(io.BytesIO(data))
        return "\n".join(page.extract_text() or "" for page in reader.pages).strip()

    if filename.endswith(".docx"):
        from docx import Document

        document = Document(io.BytesIO(data))
        paragraphs = [p.text for p in document.paragraphs if p.text.strip()]
        return "\n".join(paragraphs).strip()

    raise ValueError("Unsupported file type.")


def docx_bytes(title, content):
    from docx import Document

    document = Document()
    document.add_heading(title, 0)

    for line in content.splitlines():
        line = line.strip()

        if not line:
            continue
        elif line.startswith("### "):
            document.add_heading(line[4:], level=3)
        elif line.startswith("## "):
            document.add_heading(line[3:], level=2)
        elif line.startswith("# "):
            document.add_heading(line[2:], level=1)
        elif line.startswith(("- ", "* ")):
            document.add_paragraph(line[2:], style="List Bullet")
        elif re.match(r"^\d+[.)]\s", line):
            document.add_paragraph(
                re.sub(r"^\d+[.)]\s", "", line),
                style="List Number",
            )
        else:
            document.add_paragraph(line)

    buffer = io.BytesIO()
    document.save(buffer)
    buffer.seek(0)
    return buffer.getvalue()


# --------------------- OFFLINE FLASHCARDS -----------------------

def offline_flashcards(source, count):
    sentences = split_sentences(source)
    terms = extract_terms(source, 60)
    cards = []

    for sentence in sentences:
        words = re.findall(r"\b[A-Za-z][A-Za-z0-9_-]{3,}\b", sentence)
        candidates = [w for w in words if w.lower() in terms]

        if not candidates:
            continue

        target = max(candidates, key=len)
        masked = re.sub(
            r"\b" + re.escape(target) + r"\b",
            "________",
            sentence,
            count=1,
            flags=re.IGNORECASE,
        )

        if masked != sentence:
            cards.append({
                "question": f"Complete the source sentence: {masked}",
                "answer": target,
                "difficult": False,
            })

        if len(cards) >= count:
            break

    # If no good blank can be made, use source comprehension prompts.
    if not cards:
        for i, sentence in enumerate(sentences[:count], start=1):
            cards.append({
                "question": f"What does the source state in point {i}?",
                "answer": sentence,
                "difficult": False,
            })

    return cards[:count]


def create_flashcards(source, count):
    prompt = f"""
Create up to {count} useful revision flashcards based ONLY on the source below.
Return only a JSON array. Each object must contain:
"question": a clear question,
"answer": a concise answer supported by the source.
Do not add outside facts. If the source cannot support that many cards,
return fewer cards.

SOURCE:
{source[:20000]}
"""

    answer, error = ai_request(
        prompt,
        "Return valid JSON only. Every flashcard must be supported by the source.",
    )

    parsed = parse_json(answer)

    if isinstance(parsed, list):
        cards = []
        for item in parsed:
            if (
                isinstance(item, dict)
                and str(item.get("question", "")).strip()
                and str(item.get("answer", "")).strip()
            ):
                cards.append({
                    "question": str(item["question"]).strip(),
                    "answer": str(item["answer"]).strip(),
                    "difficult": False,
                })

        if cards:
            return cards[:count], "AI-generated from your material."

    return offline_flashcards(source, count), (
        "Offline flashcards generated from the supplied text. "
        "Some cards use sentence-completion prompts."
    )


# ------------------------ OFFLINE QUIZ --------------------------

def offline_quiz(source, count):
    """
    Build sentence-completion questions using source sentences.
    Distractors are words taken from the same source, not outside facts.
    """
    sentences = split_sentences(source)
    terms = extract_terms(source, 100)

    if len(terms) < 2:
        return []

    cards = []

    for sentence in sentences:
        words = re.findall(r"\b[A-Za-z][A-Za-z0-9_-]{3,}\b", sentence)
        candidates = [w for w in words if w.lower() in terms]

        if not candidates:
            continue

        answer = max(candidates, key=len)
        question_text = re.sub(
            r"\b" + re.escape(answer) + r"\b",
            "________",
            sentence,
            count=1,
            flags=re.IGNORECASE,
        )

        distractors = [
            term for term in terms
            if term.lower() != answer.lower()
        ]
        random.shuffle(distractors)

        options = [answer] + distractors[:3]
        options = list(dict.fromkeys(options))

        if len(options) < 2:
            continue

        while len(options) < 4:
            options.append("Not stated in this sentence")

        options = options[:4]
        random.shuffle(options)

        cards.append({
            "question": f"Complete the sentence from your source: {question_text}",
            "options": options,
            "answer": answer,
            "explanation": (
                f"The source sentence says: {sentence}"
            ),
        })

        if len(cards) >= count:
            break

    return cards


def generate_quiz(source, count):
    prompt = f"""
Create up to {count} multiple-choice revision questions based ONLY on the
source text below. Return a valid JSON array. Each item must have:
"question": string,
"options": exactly four different strings,
"answer": exactly one option string,
"explanation": a brief explanation supported by the source.

Do not invent facts or use outside knowledge. Omit questions that cannot
be answered from the source.

SOURCE:
{source[:24000]}
"""

    answer, error = ai_request(
        prompt,
        "You create source-grounded quizzes. Return JSON only.",
    )

    parsed = parse_json(answer)
    valid = []

    if isinstance(parsed, list):
        for item in parsed:
            if not isinstance(item, dict):
                continue

            options = item.get("options")
            correct = str(item.get("answer", "")).strip()

            if (
                isinstance(options, list)
                and len(options) == 4
                and all(str(option).strip() for option in options)
                and correct in [str(option).strip() for option in options]
                and str(item.get("question", "")).strip()
                and str(item.get("explanation", "")).strip()
            ):
                valid.append({
                    "question": str(item["question"]).strip(),
                    "options": [str(option).strip() for option in options],
                    "answer": correct,
                    "explanation": str(item["explanation"]).strip(),
                })

    if valid:
        return valid[:count], "AI-generated from your supplied material."

    offline = offline_quiz(source, count)

    if offline:
        return offline, (
            "Offline sentence-completion quiz. Questions and explanations "
            "are based on sentences and terms from your source."
        )

    return [], (
        "Not enough clear sentences or terms were found to build a quiz. "
        "Paste more chapter material."
    )


# ---------------------- DOCUMENT Q&A ----------------------------

def offline_document_answer(source, question):
    """Return matching sentences; do not fabricate a missing answer."""
    question_words = {
        word.lower()
        for word in re.findall(r"\b[A-Za-z]{3,}\b", question)
    }

    stop = {
        "what", "when", "where", "which", "who", "whom", "does",
        "this", "that", "from", "with", "about", "explain", "tell",
        "give", "please", "describe", "into", "your", "have",
    }
    question_words -= stop

    sentences = split_sentences(source)
    ranked = []

    for sentence in sentences:
        sentence_words = {
            word.lower()
            for word in re.findall(r"\b[A-Za-z]{3,}\b", sentence)
        }
        overlap = len(question_words & sentence_words)

        if overlap:
            ranked.append((overlap, sentence))

    ranked.sort(key=lambda item: item[0], reverse=True)

    if not ranked:
        return (
            "I could not find a sentence that clearly matches your question "
            "in this document. Try using keywords that appear in the file."
        )

    matches = [sentence for _, sentence in ranked[:5]]

    return (
        "### Relevant text found in your document\n\n"
        + "\n".join(f"- {sentence}" for sentence in matches)
        + "\n\n*Offline keyword matching was used. These are relevant source "
          "sentences, not a generated explanation.*"
    )


# --------------------- SPEAKAI OFFLINE FEEDBACK -----------------

def offline_language_feedback(text, language, level, goal, topic):
    words = re.findall(r"\b[\w'-]+\b", text)
    sentences = split_sentences(text)
    feedback = [
        "## Your practice review",
        "",
        f"- **Target language:** {language}",
        f"- **Practice level:** {level}",
        f"- **Goal:** {goal}",
        f"- **Prompt:** {topic}",
        f"- **Word count:** {len(words)}",
        f"- **Sentence count:** {len(sentences)}",
        "",
        "## What the offline tool can tell you",
        "- Your response has been captured as text.",
        "- Review spelling, punctuation, sentence length, and whether each sentence expresses one clear idea.",
        "- Try replacing repeated words with precise vocabulary that you know is correct.",
        "",
        "## Suggested practice",
        "1. Read your answer aloud slowly.",
        "2. Divide long sentences into shorter, clearer sentences.",
        "3. Add a specific example where it supports your point.",
        "4. Record a second attempt and compare its transcript with this one.",
        "",
        "**Limitation:** This offline check does not reliably detect all grammar "
        "errors and does not assess pronunciation, accent, or intonation. "
        "Configure an AI API key for detailed language feedback.",
    ]

    return "\n".join(feedback)


# ----------------------------- CSS ------------------------------

st.markdown("""
<style>
.stApp {
    background:
      radial-gradient(circle at 10% 0%, #19264a 0, transparent 38%),
      linear-gradient(135deg, #080d18, #0c1427);
    color: #eef3ff;
}
[data-testid="stSidebar"] {
    background: #0a1223;
    border-right: 1px solid #293958;
}
.hero {
    padding: 28px;
    border: 1px solid #34476b;
    border-radius: 22px;
    background: linear-gradient(120deg, #202b52, #111d31);
    margin-bottom: 20px;
}
.kicker {
    color: #a9a0ff;
    text-transform: uppercase;
    letter-spacing: 3px;
    font-size: 12px;
    font-weight: 800;
}
.module-card {
    padding: 18px;
    min-height: 155px;
    background: #111b30;
    border: 1px solid #2b3d5e;
    border-radius: 17px;
    margin-bottom: 10px;
}
.module-card p {
    color: #b8c5dc;
}
div.stButton > button {
    border-radius: 10px;
    border: 1px solid #43577d;
    font-weight: 650;
}
[data-testid="stMetric"] {
    background: #111b30;
    border: 1px solid #293958;
    border-radius: 14px;
    padding: 12px;
}
</style>
""", unsafe_allow_html=True)


# ------------------------ SESSION STATE -------------------------

DEFAULTS = {
    "page": "Dashboard",
    "lecture_transcript": "",
    "lecture_notes": "",
    "lecture_status": "",
    "speak_feedback": "",
    "speak_text": "",
    "flashcards": [],
    "card_index": 0,
    "revealed": False,
    "quiz": [],
    "quiz_answers": {},
    "quiz_submitted": False,
    "quiz_source_title": "Study quiz",
}

for key, value in DEFAULTS.items():
    if key not in st.session_state:
        st.session_state[key] = value


PAGES = [
    "Dashboard",
    "Lecture-to-Notes",
    "SpeakAI Coach",
    "Flashcard Lab",
    "Quiz Arena",
    "Document Study",
    "Learning Analytics",
    "Resource Library",
]


# ----------------------------- SIDEBAR --------------------------

with st.sidebar:
    st.markdown("## 🎓 EduVerse AI Lab")
    st.caption("Your personal learning workspace")
    st.divider()

    chosen_page = st.radio(
        "Navigate",
        PAGES,
        index=PAGES.index(st.session_state.page),
    )
    st.session_state.page = chosen_page

    st.divider()

    if get_api_key():
        st.success("AI key detected")
    else:
        st.info("Offline mode enabled")

    st.caption("Your local database: eduverse.db")
    st.caption("Audio transcription may require internet or local Whisper + FFmpeg.")


page = st.session_state.page


# ============================================================
# DASHBOARD
# ============================================================

if page == "Dashboard":
    st.markdown("""
    <div class="hero">
      <div class="kicker">Learn · Practice · Improve</div>
      <h1>Welcome to EduVerse AI Lab</h1>
      <p>Turn your own lectures and study material into notes,
      flashcards, quizzes, and language practice.</p>
    </div>
    """, unsafe_allow_html=True)

    activities = db.execute(
        "SELECT COUNT(*) FROM activity"
    ).fetchone()[0]
    resources = db.execute(
        "SELECT COUNT(*) FROM resources"
    ).fetchone()[0]
    quiz_rows = db.execute(
        "SELECT score FROM activity WHERE kind='Quiz' AND score IS NOT NULL"
    ).fetchall()

    average = (
        sum(row[0] for row in quiz_rows) / len(quiz_rows)
        if quiz_rows else None
    )

    a, b, c = st.columns(3)
    a.metric("Activities completed", activities)
    b.metric("Saved resources", resources)
    c.metric("Average quiz score", f"{average:.1f}%" if average is not None else "—")

    st.subheader("Choose your learning module")

    modules = [
        ("🎙️ Lecture-to-Notes", "Transcribe and summarize your actual lecture."),
        ("🗣️ SpeakAI Coach", "Practice communication and review your response."),
        ("🧠 Flashcard Lab", "Create revision cards from your own chapter."),
        ("🏆 Quiz Arena", "Test your knowledge using your own material."),
        ("📄 Document Study", "Read PDF, DOCX, TXT, MD, and CSV files."),
        ("📊 Learning Analytics", "Track quizzes and completed activities."),
    ]

    cols = st.columns(3)

    for i, (title, description) in enumerate(modules):
        with cols[i % 3]:
            st.markdown(
                f'<div class="module-card"><h3>{title}</h3>'
                f'<p>{description}</p></div>',
                unsafe_allow_html=True,
            )

            target = PAGES[i + 1] if i < 4 else (
                "Document Study" if i == 4 else "Learning Analytics"
            )

            if st.button("Open module →", key=f"module_{i}", use_container_width=True):
                st.session_state.page = target
                st.rerun()

    st.info(
        "For chapter-specific results, provide the actual transcript or document "
        "and review the extracted text before generating study material."
    )


# ============================================================
# LECTURE-TO-NOTES
# ============================================================

elif page == "Lecture-to-Notes":
    st.title("🎙️ Lecture-to-Notes Studio")
    st.write(
        "Record your lecture, review the recognized words, and generate notes "
        "from that exact text."
    )

    lecture_title = st.text_input(
        "Lecture or chapter title",
        value="",
        placeholder="Example: Dynamic Programming — All-Pairs Shortest Path",
        key="lecture_title",
    )

    input_mode = st.radio(
        "Choose input",
        ["Record audio", "Upload audio", "Paste transcript"],
        horizontal=True,
        key="lecture_mode",
    )

    audio = None
    pasted_text = ""

    if input_mode == "Record audio":
        audio = st.audio_input("Record your lecture", key="lecture_record")

    elif input_mode == "Upload audio":
        audio = st.file_uploader(
            "Upload audio",
            type=["wav", "mp3", "m4a", "ogg", "flac", "aiff"],
            key="lecture_upload",
        )

    else:
        pasted_text = st.text_area(
            "Paste the actual lecture transcript",
            height=180,
            placeholder="Paste the words spoken in your lecture...",
            key="lecture_paste",
        )

    col1, col2 = st.columns(2)

    with col1:
        if input_mode == "Paste transcript":
            if st.button("Use pasted transcript", type="primary"):
                if len(normalize(pasted_text)) < 20:
                    st.warning("Paste more of the actual lecture text.")
                else:
                    st.session_state.lecture_transcript = pasted_text.strip()
                    st.session_state.lecture_notes = ""
                    st.success("Transcript loaded. Review it below.")

        elif st.button("Transcribe lecture", type="primary"):
            if audio is None:
                st.warning("Record or upload audio first.")
            else:
                with st.spinner("Transcribing your audio..."):
                    transcript, error = transcribe_audio(audio)

                if transcript:
                    st.session_state.lecture_transcript = transcript
                    st.session_state.lecture_notes = ""
                    st.session_state.lecture_status = ""
                    log_activity("Transcription", lecture_title or "Untitled lecture")
                    st.success("Transcription completed. Review the text below.")
                else:
                    st.error(error)

    with col2:
        if st.button("Clear lecture workspace"):
            st.session_state.lecture_transcript = ""
            st.session_state.lecture_notes = ""
            st.session_state.lecture_status = ""
            st.rerun()

    if st.session_state.lecture_transcript:
        st.divider()
        st.subheader("Step 1 — Review the actual transcript")

        edited_transcript = st.text_area(
            "Correct any transcription mistakes here",
            value=st.session_state.lecture_transcript,
            height=240,
            key="lecture_editor",
        )

        if st.button("Save transcript edits"):
            st.session_state.lecture_transcript = edited_transcript.strip()
            st.session_state.lecture_notes = ""
            st.success("Your transcript edits have been saved.")

        st.caption(
            "Important: the notes generator uses this transcript, not unrelated "
            "chapters. Correct any speech-recognition errors before continuing."
        )

        st.subheader("Step 2 — Generate notes from this transcript")

        if st.button("Generate chapter-specific notes", type="primary"):
            source = edited_transcript.strip()

            if len(source) < 40:
                st.warning(
                    "The transcript is too short. Add or correct the actual lecture text."
                )
            else:
                with st.spinner("Creating notes from the reviewed transcript..."):
                    notes, status = generate_notes(
                        source,
                        lecture_title.strip() or "Lecture Notes",
                    )

                if notes:
                    st.session_state.lecture_transcript = source
                    st.session_state.lecture_notes = notes
                    st.session_state.lecture_status = status
                    log_activity(
                        "Lecture notes",
                        lecture_title.strip() or "Lecture Notes",
                    )
                else:
                    st.error(status)

        if st.session_state.lecture_notes:
            st.divider()
            st.subheader("Step 3 — Your lecture notes")
            st.caption(st.session_state.lecture_status)
            st.markdown(st.session_state.lecture_notes)

            st.download_button(
                "Download notes as Word",
                data=docx_bytes(
                    lecture_title or "Lecture Notes",
                    st.session_state.lecture_notes,
                ),
                file_name="eduverse_lecture_notes.docx",
                mime="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            )

            st.download_button(
                "Download transcript as TXT",
                data=st.session_state.lecture_transcript,
                file_name="lecture_transcript.txt",
                mime="text/plain",
            )

            if st.button("Save notes to Resource Library"):
                save_resource(
                    lecture_title or "Lecture Notes",
                    "Lecture notes",
                    st.session_state.lecture_notes,
                )
                st.success("Notes saved to Resource Library.")


# ============================================================
# SPEAKAI COACH
# ============================================================

elif page == "SpeakAI Coach":
    st.title("🗣️ SpeakAI Communication Coach")
    st.write(
        "Practice an answer, review the recognized text, and get "
        "language feedback without claiming to score pronunciation from text."
    )

    language = st.selectbox(
        "Language you are practising",
        ["English", "Hindi", "Marathi", "Spanish", "French", "German"],
    )
    level = st.selectbox(
        "Your level",
        ["Beginner", "Elementary", "Intermediate", "Advanced"],
    )
    goal = st.selectbox(
        "Practice goal",
        [
            "Everyday conversation",
            "College presentation",
            "Job interview",
            "Professional communication",
            "Grammar and vocabulary",
        ],
    )
    topic = st.text_input(
        "Practice prompt",
        value="Introduce yourself and explain your current learning goal.",
    )

    mode = st.radio(
        "Your answer",
        ["Record audio", "Upload audio", "Type answer"],
        horizontal=True,
        key="speak_mode",
    )

    speak_audio = None
    typed_answer = ""

    if mode == "Record audio":
        speak_audio = st.audio_input("Record your answer", key="speak_record")

    elif mode == "Upload audio":
        speak_audio = st.file_uploader(
            "Upload your answer",
            type=["wav", "mp3", "m4a", "ogg", "flac", "aiff"],
            key="speak_upload",
        )

    else:
        typed_answer = st.text_area(
            "Type the answer you would say aloud",
            height=150,
            placeholder="Write your answer here...",
            key="speak_typed",
        )

    if st.button("Review my answer", type="primary"):
        if mode == "Type answer":
            response_text = typed_answer.strip()
            error = None
        else:
            with st.spinner("Transcribing your answer..."):
                response_text, error = transcribe_audio(speak_audio)

        if error:
            st.error(error)

        elif not response_text or len(normalize(response_text)) < 8:
            st.warning("Please provide a longer answer so the review is useful.")

        else:
            st.session_state.speak_text = response_text

            prompt = f"""
You are a supportive language tutor.

Language being practised: {language}
Learner level: {level}
Goal: {goal}
Practice prompt: {topic}
Learner's exact response: {response_text}

Give specific, honest feedback with these sections:
1. What was communicated clearly
2. Definite grammar or sentence issues, quoting short exact fragments
3. Suggested corrections and why they are better
4. Three useful vocabulary or phrase improvements
5. A sample improved answer that does not invent personal facts
6. One short practice exercise

If you cannot establish that something is wrong, call it a style suggestion,
not a definite error. If the response does not match the target language,
mention that. Do not claim to assess pronunciation from text.
"""

            with st.spinner("Reviewing your answer..."):
                feedback, error = ai_request(
                    prompt,
                    "You are an accurate, kind language coach. Never invent learner mistakes.",
                )

            st.subheader("Recognized answer")
            st.write(response_text)

            if feedback:
                st.subheader("Personalized feedback")
                st.markdown(feedback)
                st.session_state.speak_feedback = feedback

            else:
                feedback = offline_language_feedback(
                    response_text, language, level, goal, topic
                )
                st.subheader("Offline review")
                st.markdown(feedback)
                st.info(
                    "Detailed grammar corrections require an AI API key. "
                    "The offline review reports text statistics and practice guidance."
                )
                st.session_state.speak_feedback = feedback

            log_activity("SpeakAI practice", f"{language} — {goal}")

            if st.button("Save feedback to Resource Library"):
                save_resource(
                    f"{language} practice feedback",
                    "SpeakAI feedback",
                    f"Response:\n{response_text}\n\n{st.session_state.speak_feedback}",
                )
                st.success("Feedback saved.")

    st.caption(
        "Text transcription cannot reliably assess pronunciation, accent, "
        "intonation, or spoken fluency."
    )


# ============================================================
# FLASHCARD LAB
# ============================================================

elif page == "Flashcard Lab":
    st.title("🧠 Flashcard Lab")
    st.write("Generate revision cards from your own chapter or lecture.")

    source = st.text_area(
        "Paste chapter notes or transcript",
        height=190,
        key="flashcard_source",
        placeholder="Paste the material you want to revise...",
    )

    count = st.slider("Number of flashcards", 3, 20, 8)

    if st.button("Create flashcards", type="primary"):
        if len(normalize(source)) < 50:
            st.warning("Paste more meaningful chapter material first.")
        else:
            with st.spinner("Creating flashcards..."):
                cards, status = create_flashcards(source, count)

            if cards:
                st.session_state.flashcards = cards
                st.session_state.card_index = 0
                st.session_state.revealed = False
                st.success(status)
                log_activity("Flashcards", "Source-based flashcards")
            else:
                st.error(status)

    cards = st.session_state.flashcards

    if cards:
        st.divider()
        st.caption(f"{len(cards)} cards available")

        idx = min(st.session_state.card_index, len(cards) - 1)
        card = cards[idx]

        st.markdown(f"### Card {idx + 1} of {len(cards)}")
        st.markdown(
            f'<div class="module-card"><h3>Question</h3>'
            f'<p>{card["question"]}</p></div>',
            unsafe_allow_html=True,
        )

        if st.session_state.revealed:
            st.markdown("#### Answer")
            st.info(card["answer"])

        if st.button(
            "Hide answer" if st.session_state.revealed else "Reveal answer"
        ):
            st.session_state.revealed = not st.session_state.revealed
            st.rerun()

        left, middle, right = st.columns(3)

        with left:
            if st.button("← Previous", disabled=idx == 0):
                st.session_state.card_index -= 1
                st.session_state.revealed = False
                st.rerun()

        with middle:
            if st.button("Mark difficult"):
                cards[idx]["difficult"] = True
                st.session_state.flashcards = cards
                st.success("Marked for review.")

        with right:
            if st.button("Next →", disabled=idx >= len(cards) - 1):
                st.session_state.card_index += 1
                st.session_state.revealed = False
                st.rerun()

        difficult_cards = [c for c in cards if c.get("difficult")]

        if difficult_cards and st.button(
            f"Review difficult cards ({len(difficult_cards)})"
        ):
            st.session_state.flashcards = difficult_cards
            st.session_state.card_index = 0
            st.session_state.revealed = False
            st.rerun()

        export_text = "\n\n".join(
            f"Q: {card['question']}\nA: {card['answer']}" for card in cards
        )

        st.download_button(
            "Download flashcards as TXT",
            data=export_text,
            file_name="eduverse_flashcards.txt",
            mime="text/plain",
        )


# ============================================================
# QUIZ ARENA
# ============================================================

elif page == "Quiz Arena":
    st.title("🏆 Quiz Arena")
    st.write("Build a quiz from your actual chapter notes or lecture transcript.")

    quiz_title = st.text_input("Chapter name", value="My chapter")
    source = st.text_area(
        "Paste the chapter material",
        height=190,
        key="quiz_source",
        placeholder="Paste the source text for your quiz...",
    )
    number = st.slider("Number of questions", 3, 15, 5)

    if st.button("Generate quiz", type="primary"):
        if len(normalize(source)) < 80:
            st.warning("Paste more chapter content before generating a quiz.")
        else:
            with st.spinner("Creating source-based questions..."):
                quiz, status = generate_quiz(source, number)

            if quiz:
                st.session_state.quiz = quiz
                st.session_state.quiz_answers = {}
                st.session_state.quiz_submitted = False
                st.session_state.quiz_source_title = quiz_title.strip() or "Study quiz"
                st.success(status)
            else:
                st.error(status)

    quiz = st.session_state.quiz

    if quiz:
        st.divider()
        st.subheader(st.session_state.quiz_source_title)

        if not st.session_state.quiz_submitted:
            with st.form("eduverse_quiz_form"):
                for i, question in enumerate(quiz):
                    st.markdown(f"**Question {i + 1}. {question['question']}**")

                    choice = st.radio(
                        "Select one answer",
                        question["options"],
                        index=None,
                        key=f"answer_{i}",
                        label_visibility="collapsed",
                    )

                    st.session_state.quiz_answers[str(i)] = choice

                submitted = st.form_submit_button("Submit quiz")

            if submitted:
                score = sum(
                    st.session_state.quiz_answers.get(str(i)) == q["answer"]
                    for i, q in enumerate(quiz)
                )
                percentage = 100 * score / len(quiz)

                st.session_state.quiz_submitted = True
                log_activity(
                    "Quiz",
                    st.session_state.quiz_source_title,
                    percentage,
                )
                st.rerun()

        if st.session_state.quiz_submitted:
            score = sum(
                st.session_state.quiz_answers.get(str(i)) == q["answer"]
                for i, q in enumerate(quiz)
            )
            percentage = 100 * score / len(quiz)

            st.metric("Your score", f"{score} / {len(quiz)}", f"{percentage:.0f}%")

            for i, question in enumerate(quiz):
                selected = st.session_state.quiz_answers.get(str(i))

                if selected == question["answer"]:
                    st.success(f"Question {i + 1}: Correct")
                else:
                    st.error(
                        f"Question {i + 1}: Incorrect. "
                        f"Your answer: {selected or 'Not answered'}. "
                        f"Correct answer: {question['answer']}"
                    )

                st.caption(question["explanation"])

            if st.button("Start a new quiz"):
                st.session_state.quiz = []
                st.session_state.quiz_answers = {}
                st.session_state.quiz_submitted = False
                st.rerun()


# ============================================================
# DOCUMENT STUDY ASSISTANT
# ============================================================

elif page == "Document Study":
    st.title("📄 Document Study Assistant")
    st.write(
        "Upload a study document, inspect the extracted text, create notes, "
        "or ask questions about its content."
    )

    upload = st.file_uploader(
        "Upload PDF, DOCX, TXT, MD, or CSV",
        type=["pdf", "docx", "txt", "md", "csv"],
    )

    if upload:
        try:
            extracted = extract_document(upload)
        except Exception as exc:
            extracted = ""
            st.error(f"Could not read the document: {exc}")

        if not extracted:
            st.warning(
                "No readable text was found. A scanned/image-only PDF may need OCR."
            )

        else:
            st.success(
                f"Extracted {len(extracted):,} characters from {upload.name}."
            )

            with st.expander("Review extracted document text"):
                st.text_area(
                    "Extracted text",
                    value=extracted,
                    height=220,
                    key="extracted_document_preview",
                )

            question = st.text_input(
                "Ask a question about this document",
                placeholder="Example: What are the main points?",
            )

            left, right = st.columns(2)

            with left:
                if st.button("Answer from document", type="primary"):
                    if not question.strip():
                        st.warning("Enter a question first.")
                    else:
                        prompt = f"""
Answer the user's question using ONLY the document text below.
If the answer is absent, clearly say the document does not provide enough
information. Quote relevant supporting sentences where possible.

QUESTION:
{question}

DOCUMENT:
{extracted[:30000]}
"""

                        with st.spinner("Finding a document-grounded answer..."):
                            answer, error = ai_request(
                                prompt,
                                "You answer questions from supplied documents only.",
                            )

                        if answer:
                            st.session_state["last_doc_answer"] = answer
                            st.markdown(answer)
                            log_activity("Document Q&A", upload.name)
                        else:
                            answer = offline_document_answer(extracted, question)
                            st.session_state["last_doc_answer"] = answer
                            st.markdown(answer)
                            st.info(
                                "Offline keyword matching was used. "
                                "For a generated explanation, configure an AI API key."
                            )

            with right:
                if st.button("Create notes from document"):
                    with st.spinner("Creating document notes..."):
                        notes, status = generate_notes(extracted, upload.name)

                    st.session_state["last_doc_notes"] = notes
                    st.session_state["last_doc_notes_status"] = status
                    log_activity("Document notes", upload.name)

            if st.session_state.get("last_doc_notes"):
                st.divider()
                st.subheader("Document notes")
                st.caption(st.session_state.get("last_doc_notes_status", ""))
                st.markdown(st.session_state["last_doc_notes"])

                st.download_button(
                    "Download document notes as Word",
                    data=docx_bytes(
                        upload.name,
                        st.session_state["last_doc_notes"],
                    ),
                    file_name="eduverse_document_notes.docx",
                    mime="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
                )

            if st.session_state.get("last_doc_answer"):
                if st.button("Save latest answer to Resource Library"):
                    save_resource(
                        f"Answer: {question[:100]}",
                        "Document answer",
                        st.session_state["last_doc_answer"],
                    )
                    st.success("Answer saved.")


# ============================================================
# LEARNING ANALYTICS
# ============================================================

elif page == "Learning Analytics":
    st.title("📊 Learning Analytics")

    rows = db.execute("""
        SELECT kind, title, score, created
        FROM activity
        ORDER BY id DESC
        LIMIT 200
    """).fetchall()

    total = db.execute("SELECT COUNT(*) FROM activity").fetchone()[0]
    quiz_rows = db.execute(
        "SELECT score FROM activity WHERE kind='Quiz' AND score IS NOT NULL"
    ).fetchall()

    average = (
        sum(row[0] for row in quiz_rows) / len(quiz_rows)
        if quiz_rows else None
    )

    a, b, c = st.columns(3)
    a.metric("Total activities", total)
    b.metric("Quizzes completed", len(quiz_rows))
    c.metric(
        "Average quiz score",
        f"{average:.1f}%" if average is not None else "No quizzes yet",
    )

    if rows:
        st.dataframe(
            [
                {
                    "Activity": row[0],
                    "Title": row[1],
                    "Score": f"{row[2]:.1f}%" if row[2] is not None else "—",
                    "Date": row[3].replace("T", " "),
                }
                for row in rows
            ],
            use_container_width=True,
            hide_index=True,
        )

        if st.button("Export activity history as CSV"):
            csv_lines = ["Activity,Title,Score,Date"]
            for row in rows:
                values = [
                    str(row[0]), str(row[1]),
                    "" if row[2] is None else str(row[2]),
                    str(row[3]),
                ]
                safe = [
                    '"' + value.replace('"', '""') + '"'
                    for value in values
                ]
                csv_lines.append(",".join(safe))

            st.download_button(
                "Download activity CSV",
                data="\n".join(csv_lines),
                file_name="eduverse_activity.csv",
                mime="text/csv",
            )
    else:
        st.info("Your activities will appear here after you use the modules.")


# ============================================================
# RESOURCE LIBRARY
# ============================================================

elif page == "Resource Library":
    st.title("📚 Resource Library")
    st.write("Save notes, revision material, and language-practice feedback locally.")

    with st.form("manual_resource_form"):
        title = st.text_input("Resource title")
        body = st.text_area("Resource content", height=150)
        submitted = st.form_submit_button("Save resource")

    if submitted:
        if not title.strip() or not body.strip():
            st.warning("Enter both a title and content.")
        else:
            save_resource(title.strip(), "Personal note", body.strip())
            st.success("Resource saved.")
            st.rerun()

    saved = db.execute("""
        SELECT id, title, kind, body, created
        FROM resources
        ORDER BY id DESC
    """).fetchall()

    if not saved:
        st.info("No resources saved yet.")

    for rid, title, kind, body, created in saved:
        with st.expander(f"{title} · {kind} · {created[:10]}"):
            st.markdown(body)

            st.download_button(
                "Download resource as TXT",
                data=body,
                file_name=f"eduverse_resource_{rid}.txt",
                mime="text/plain",
                key=f"resource_download_{rid}",
            )

            if st.button("Delete resource", key=f"resource_delete_{rid}"):
                db.execute("DELETE FROM resources WHERE id=?", (rid,))
                db.commit()
                st.success("Resource deleted.")
                st.rerun()


# ------------------------------ FOOTER ---------------------------

st.divider()
st.caption(
    "EduVerse AI Lab · Study support only. Always verify important notes "
    "against your original lecture or course material."
)
